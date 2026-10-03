"""
Reporter module for Business Sales & Performance Analytics System.
Generates an automated, formatted Excel business summary report (.xlsx) using openpyxl.
"""

import logging
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("SalesAnalyticsPipeline")


def generate_excel_report(
    kpis: dict,
    branch_df: pd.DataFrame,
    top_prod_df: pd.DataFrame,
    val_report: dict,
    anomalies_df: pd.DataFrame,
    forecast_df: pd.DataFrame,
    output_path: Path
):
    """
    Generates multi-tab Excel business performance summary report.
    
    :param kpis: KPI dictionary.
    :param branch_df: Branch performance DataFrame.
    :param top_prod_df: Top products DataFrame.
    :param val_report: Data validation summary dictionary.
    :param anomalies_df: Detected anomalies DataFrame.
    :param forecast_df: Forecast DataFrame.
    :param output_path: Destination file path for .xlsx report.
    """
    logger.info(f"Generating formatted Excel summary report at: {output_path}...")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default active sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    section_font = Font(name="Calibri", size=12, bold=True, color="2F5597")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # -------------------------------------------------------------
    # SHEET 1: Executive KPI Summary & Forecast
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Executive Summary")
    ws1.cell(row=1, column=1, value="Business Executive KPI Summary").font = title_font
    
    kpi_rows = [
        ("Total Net Revenue", f"${kpis['total_revenue']:,.2f}"),
        ("Total Transactions", f"{kpis['total_transactions']:,}"),
        ("Total Units Sold", f"{kpis['total_units_sold']:,}"),
        ("Average Order Value (AOV)", f"${kpis['average_order_value']:,.2f}"),
        ("Average Discount Rate", f"{kpis['average_discount_pct']:.1f}%")
    ]

    ws1.cell(row=3, column=1, value="Core Business Metrics").font = section_font
    ws1.cell(row=4, column=1, value="KPI Metric").font = header_font
    ws1.cell(row=4, column=1).fill = header_fill
    ws1.cell(row=4, column=2, value="Value").font = header_font
    ws1.cell(row=4, column=2).fill = header_fill

    for r_idx, (metric, val) in enumerate(kpi_rows, start=5):
        c1 = ws1.cell(row=r_idx, column=1, value=metric)
        c2 = ws1.cell(row=r_idx, column=2, value=val)
        c1.border = thin_border
        c2.border = thin_border

    # Forecast table below KPIs
    ws1.cell(row=12, column=1, value="Monthly Revenue & 3-Month Forecast").font = section_font
    ws1.cell(row=13, column=1, value="Month").font = header_font
    ws1.cell(row=13, column=1).fill = header_fill
    ws1.cell(row=13, column=2, value="Actual Revenue ($)").font = header_font
    ws1.cell(row=13, column=2).fill = header_fill
    ws1.cell(row=13, column=3, value="Linear Trend ($)").font = header_font
    ws1.cell(row=13, column=3).fill = header_fill
    ws1.cell(row=13, column=4, value="3M Moving Avg ($)").font = header_font
    ws1.cell(row=13, column=4).fill = header_fill
    ws1.cell(row=13, column=5, value="Status").font = header_font
    ws1.cell(row=13, column=5).fill = header_fill

    for row_idx, row in enumerate(forecast_df.to_dict('records'), start=14):
        c1 = ws1.cell(row=row_idx, column=1, value=row["year_month"])
        c2 = ws1.cell(row=row_idx, column=2, value=row["total_revenue"] if pd.notna(row["total_revenue"]) else "-")
        c3 = ws1.cell(row=row_idx, column=3, value=row["linear_trend"])
        c4 = ws1.cell(row=row_idx, column=4, value=row["moving_avg_3m"])
        c5 = ws1.cell(row=row_idx, column=5, value="Forecast" if row["is_forecast"] else "Actual")

        for cell in (c1, c2, c3, c4, c5):
            cell.border = thin_border
            if row["is_forecast"]:
                cell.font = Font(bold=True, color="C00000")

    # -------------------------------------------------------------
    # SHEET 2: Branch & Product Breakdown
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Sales Breakdown")
    ws2.cell(row=1, column=1, value="Branch Performance Breakdown").font = title_font

    headers_branch = ["Region", "Branch ID", "Branch Name", "Total Revenue ($)", "Transactions", "Units Sold", "AOV ($)"]
    for c_idx, h in enumerate(headers_branch, start=1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for r_idx, row in enumerate(branch_df.to_dict('records'), start=4):
        ws2.cell(row=r_idx, column=1, value=row["region"]).border = thin_border
        ws2.cell(row=r_idx, column=2, value=row["branch_id"]).border = thin_border
        ws2.cell(row=r_idx, column=3, value=row["branch_name"]).border = thin_border
        ws2.cell(row=r_idx, column=4, value=row["total_revenue"]).border = thin_border
        ws2.cell(row=r_idx, column=5, value=row["total_transactions"]).border = thin_border
        ws2.cell(row=r_idx, column=6, value=row["units_sold"]).border = thin_border
        ws2.cell(row=r_idx, column=7, value=row["avg_order_value"]).border = thin_border

    start_prod_row = 4 + len(branch_df) + 3
    ws2.cell(row=start_prod_row, column=1, value="Top Performing Products").font = section_font
    
    headers_prod = ["Product ID", "Product Name", "Category", "Total Revenue ($)", "Units Sold", "Avg Unit Price ($)"]
    for c_idx, h in enumerate(headers_prod, start=1):
        cell = ws2.cell(row=start_prod_row + 1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for r_idx, row in enumerate(top_prod_df.to_dict('records'), start=start_prod_row + 2):
        ws2.cell(row=r_idx, column=1, value=row["product_id"]).border = thin_border
        ws2.cell(row=r_idx, column=2, value=row["product_name"]).border = thin_border
        ws2.cell(row=r_idx, column=3, value=row["category"]).border = thin_border
        ws2.cell(row=r_idx, column=4, value=row["total_revenue"]).border = thin_border
        ws2.cell(row=r_idx, column=5, value=row["units_sold"]).border = thin_border
        ws2.cell(row=r_idx, column=6, value=row["avg_unit_price"]).border = thin_border

    # -------------------------------------------------------------
    # SHEET 3: Data Quality & Anomalies Report
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Data Quality & Anomalies")
    ws3.cell(row=1, column=1, value="Data Quality Audit & Anomaly Detection").font = title_font

    ws3.cell(row=3, column=1, value="Raw File Validation Audit").font = section_font
    val_rows = [
        ("Total Records Processed", val_report["total_records"]),
        ("Valid Records Retained", val_report["valid_records"]),
        ("Flawed/Invalid Records Dropped", val_report["invalid_records"]),
        ("Exact Duplicate Rows Removed", val_report["exact_duplicates"]),
        ("Duplicate Transaction IDs Removed", val_report["duplicate_tx_ids"]),
        ("Invalid Quantity Records", val_report["invalid_quantities"]),
        ("Invalid Price Records", val_report["invalid_unit_prices"]),
        ("Malformed Date Records", val_report["malformed_dates"])
    ]

    ws3.cell(row=4, column=1, value="Audit Metric").font = header_font
    ws3.cell(row=4, column=1).fill = header_fill
    ws3.cell(row=4, column=2, value="Count").font = header_font
    ws3.cell(row=4, column=2).fill = header_fill

    for r_idx, (metric, count) in enumerate(val_rows, start=5):
        c1 = ws3.cell(row=r_idx, column=1, value=metric)
        c2 = ws3.cell(row=r_idx, column=2, value=count)
        c1.border = thin_border
        c2.border = thin_border

    start_anom_row = 5 + len(val_rows) + 2
    ws3.cell(row=start_anom_row, column=1, value="Top Flagged High-Value Anomalous Transactions (IQR Method)").font = section_font

    headers_anom = ["TXN ID", "Date", "Branch", "Product", "Quantity", "Net Sales ($)", "Anomaly Flag"]
    for c_idx, h in enumerate(headers_anom, start=1):
        cell = ws3.cell(row=start_anom_row + 1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    sample_anom = anomalies_df.head(15).to_dict('records') if not anomalies_df.empty else []
    for r_idx, row in enumerate(sample_anom, start=start_anom_row + 2):
        ws3.cell(row=r_idx, column=1, value=row["transaction_id"]).border = thin_border
        ws3.cell(row=r_idx, column=2, value=str(row["transaction_date"])).border = thin_border
        ws3.cell(row=r_idx, column=3, value=row["branch_name"]).border = thin_border
        ws3.cell(row=r_idx, column=4, value=row["product_name"]).border = thin_border
        ws3.cell(row=r_idx, column=5, value=row["quantity"]).border = thin_border
        ws3.cell(row=r_idx, column=6, value=row["net_sales"]).border = thin_border
        ws3.cell(row=r_idx, column=7, value=row.get("anomaly_type", "High Outlier")).border = thin_border

    # Adjust Column Widths automatically across sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)
    logger.info(f"Excel report saved successfully.")
